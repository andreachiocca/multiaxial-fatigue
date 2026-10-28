"""Build small synthetic workbooks that follow the input contract (docs/workbook-format.md)."""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook

from multifat.imp.schema import material_id

SUMMARY_HEADERS = [
    "Article DOI",
    "Authors",
    "Article year",
    "Material name",
    "Young's modulus (MPa)",
    "Poisson's ratio (-)",
    "Tensile strength (MPa)",
    "Yield strength (MPa)",
    "Tensile fatigue limit* (MPa)",
    "Torsional fatigue limit* (MPa)",
    "k FS (-)",
    "k FI (-)",
]
SUMMARY_VALUES = [None, None, None, "Synthetic steel", 206000, 0.3, 900, 750, 420, 300, 1.0, 0.3]

DATASET_HEADERS = [
    "Number of cycles to failure",
    "Run out (Yes/No)",
    "Notched",
    "Phase radial load (deg)",
    "Phase torsional load (deg)",
    "Mean RADIAL stress  (MPa)",
    "Amplitute RADIAL stress  (MPa)",
    "Mean AXIAL stress  (MPa)",
    "Amplitute AXIAL stress  (MPa)",
    "Mean SHEAR stress  (MPa)",
    "Amplitute SHEAR stress  (MPa)",
]


def basquin_rows(sf=1500.0, b=-0.09, tf=860.0, b0=-0.07):
    """Uniaxial and torsion R=-1 rows on exact Basquin lines, plus multiaxial and filtered rows."""
    rows = []
    for N in (1e4, 3e4, 1e5, 3e5, 1e6, 2e6):
        rows.append([N, "No", "No", 0, 0, 0, 0, 0, sf * (2 * N) ** b, 0, 0])
    for N in (2e4, 6e4, 2e5, 6e5, 1.5e6):
        rows.append([N, "No", "No", 0, 0, 0, 0, 0, 0, 0, tf * (2 * N) ** b0])
    for N, phase in ((5e4, 0), (2e5, 0), (8e4, 90), (4e5, 90)):
        s = 0.5 * sf * (2 * N) ** b
        rows.append([N, "No", "No", 0, phase, 0, 0, 0, s, 0, s / 3**0.5])
    rows.append(
        [5e6, "Yes", "No", 0, 0, 0, 0, 0, 300, 0, 0]
    )  # run-out: kept apart from the failures
    rows.append([5e4, "No", "Yes", 0, 0, 0, 0, 0, 300, 0, 0])  # notched: dropped
    return rows


def write_workbook(
    path: Path,
    *,
    summary_headers=None,
    summary_values=None,
    dataset_headers=None,
    rows=None,
    extra_summary: dict | None = None,
    sheets=("Summary", "Dataset"),
) -> Path:
    sh = list(summary_headers or SUMMARY_HEADERS)
    sv = list(summary_values or SUMMARY_VALUES)
    for k, v in (extra_summary or {}).items():
        sh.append(k)
        sv.append(v)
    wb = Workbook()
    wb.remove(wb.active)
    if "Summary" in sheets:
        ws = wb.create_sheet("Summary")
        ws.append(sh)
        ws.append(sv)
        ws.append(["Average"] + sv[1:])
    if "Dataset" in sheets:
        ws = wb.create_sheet("Dataset")
        ws.append(list(dataset_headers or DATASET_HEADERS))
        for r in rows if rows is not None else basquin_rows():
            ws.append(r)
    wb.save(path)
    return path


# Revised format (October 2026 collection): explicit axial phase, frequencies, pressure, specimen.
REVISED_SUMMARY_HEADERS = [
    "Article DOI / Reference",
    "Authors",
    "Article year",
    "Material name",
    "Young's modulus (MPa)",
    "Poisson's ratio (-)",
    "Tensile strength (MPa)",
    "Yield strength (MPa)",
    "Tensile fatigue limit (MPa)",
    "Torsional fatigue limit (MPa)",
    "Number of Cycles (cycles)",
    "Stress ratio R",
]
REVISED_DATASET_HEADERS = [
    "Specimen name",
    "Number of cycles to failure",
    "Run out (Yes/No)",
    "Notched",
    "Phase RADIAL load (°)",
    "Phase AXIAL load (°)",
    "Phase TORSIONAL load (°)",
    "Freq RADIAL load [Hz]",
    "Freq AXIAL load[Hz]",
    "Freq TORSIONAL load [Hz]",
    "Mean RADIAL stress  (MPa)",
    "Amplitute RADIAL stress  (MPa)",
    "Mean AXIAL stress  (MPa)",
    "Amplitute AXIAL stress  (MPa)",
    "derived from (AXIAL/BENDING) load",
    "Mean SHEAR stress  (MPa)",
    "Amplitute SHEAR stress  (MPa)",
    "Mean internal PRESSURE  (MPa)",
    "Amplitute internal PRESSURE (MPa)",
    "Mean AXIAL load (N)",
    "Amplitute AXIAL load (N)",
]


def revised_row(overrides: dict | None = None) -> list:
    """One revised Dataset row: a plain failure at 1e5 cycles, axial 300 MPa, unless overridden."""
    base = {
        "Specimen name": "S1",
        "Number of cycles to failure": 1e5,
        "Run out (Yes/No)": "No",
        "Notched": "No",
        "Phase RADIAL load (°)": 0,
        "Phase AXIAL load (°)": 0,
        "Phase TORSIONAL load (°)": 0,
        "Freq RADIAL load [Hz]": "Null",
        "Freq AXIAL load[Hz]": 20,
        "Freq TORSIONAL load [Hz]": 20,
        "Mean RADIAL stress  (MPa)": 0,
        "Amplitute RADIAL stress  (MPa)": 0,
        "Mean AXIAL stress  (MPa)": 0,
        "Amplitute AXIAL stress  (MPa)": 300,
        "derived from (AXIAL/BENDING) load": "AXIAL",
        "Mean SHEAR stress  (MPa)": 0,
        "Amplitute SHEAR stress  (MPa)": 0,
        "Mean internal PRESSURE  (MPa)": 0,
        "Amplitute internal PRESSURE (MPa)": 0,
        "Mean AXIAL load (N)": 0,
        "Amplitute AXIAL load (N)": 9000,
    }
    base.update(overrides or {})
    unknown = set(base) - set(REVISED_DATASET_HEADERS)
    assert not unknown, unknown
    return [base[h] for h in REVISED_DATASET_HEADERS]


def write_revised_workbook(path: Path, rows: list[list], *, summary_rows=None) -> Path:
    """Revised-format workbook; `summary_rows` are source rows (the Average is their copy)."""
    rows_s = summary_rows or [
        ["ref", "A", 2020, material_id(path.stem), 206000, 0.3, 900, 750, 420, 300, 2e6, -1]
    ]
    wb = Workbook()
    wb.remove(wb.active)
    ws = wb.create_sheet("Summary")
    ws.append(REVISED_SUMMARY_HEADERS)
    for r in rows_s:
        ws.append(r)
    avg = []
    for j in range(4, len(REVISED_SUMMARY_HEADERS)):
        vals = [r[j] for r in rows_s if isinstance(r[j], (int, float))]
        avg.append(sum(vals) / len(vals) if vals else "Null")
    ws.append(["Average", None, None, None] + avg)
    ds = wb.create_sheet("Dataset")
    ds.append(REVISED_DATASET_HEADERS)
    for r in rows:
        ds.append(r)
    wb.save(path)
    return path
