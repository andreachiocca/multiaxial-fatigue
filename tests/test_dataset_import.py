"""Revised workbook format (October 2026 collection): material IDs, phases, loading issues.

Ported from upstream tests/test_dataset_import.py (A. Chiocca, 904c619) to pytest and our API.
"""

from pathlib import Path

import numpy as np
import pytest
from openpyxl import load_workbook

from multifat.imp.experimental_data import (
    import_experimental_data_grouped_by_material,
    load_material_workbook,
)
from multifat.imp.schema import WorkbookError, ValidationReport, material_id
from workbook_factory import revised_row, write_revised_workbook

ROOT = Path(__file__).resolve().parents[1]


def _tension_rows():
    return [
        revised_row(
            {
                "Specimen name": f"T{i}",
                "Number of cycles to failure": n,
                "Amplitute AXIAL stress  (MPa)": s,
            }
        )
        for i, (n, s) in enumerate([(1e4, 600), (1e5, 450), (1e6, 350)])
    ]


def _one(path):
    (block,) = load_material_workbook(path).values()
    return block


def test_prefix_is_not_identity(tmp_path):
    from multifat.cli.run import resolve_workbook

    assert material_id("6_42CrMo4_QT") == "42CrMo4_QT"
    assert material_id("Ti6Al4V_andrea_(AXIAL)") == "Ti6Al4V_andrea_(AXIAL)"
    write_revised_workbook(tmp_path / "6_Synth.xlsx", _tension_rows())
    assert list(import_experimental_data_grouped_by_material(tmp_path)) == ["Synth"]
    assert resolve_workbook("Synth", tmp_path) == tmp_path / "6_Synth.xlsx"
    write_revised_workbook(tmp_path / "Synth.xlsx", _tension_rows())
    with pytest.raises(WorkbookError, match="Synth"):
        resolve_workbook("Synth", tmp_path)
    with pytest.raises(WorkbookError, match="Synth"):
        import_experimental_data_grouped_by_material(tmp_path)


def test_archive_not_imported(tmp_path):
    write_revised_workbook(tmp_path / "A.xlsx", _tension_rows())
    (tmp_path / "archive").mkdir()
    write_revised_workbook(tmp_path / "archive" / "A.xlsx", _tension_rows())
    assert list(import_experimental_data_grouped_by_material(tmp_path)) == ["A"]


def test_axial_phase_is_reference(tmp_path):
    row = revised_row(
        {
            "Phase AXIAL load (°)": 30,
            "Phase TORSIONAL load (°)": 120,
            "Amplitute SHEAR stress  (MPa)": 100,
        }
    )
    block = _one(write_revised_workbook(tmp_path / "P.xlsx", _tension_rows() + [row]))
    case = block["cases"][3]
    assert case["meta"]["phase_tor_deg"] == 90
    assert case["meta"]["phase_rad_deg"] == -30
    assert block["source"]["format"] == "revised"


def test_frequency_text_kept(tmp_path):
    ranged = revised_row({"Freq AXIAL load[Hz]": "15-20", "Freq TORSIONAL load [Hz]": "15-20"})
    mixed = revised_row(
        {
            "Specimen name": "M",
            "Freq AXIAL load[Hz]": 10,
            "Freq TORSIONAL load [Hz]": 20,
            "Amplitute SHEAR stress  (MPa)": 100,
        }
    )
    rep = ValidationReport(file="F.xlsx")
    path = write_revised_workbook(tmp_path / "F.xlsx", _tension_rows() + [ranged, mixed])
    (block,) = load_material_workbook(path, report=rep).values()
    assert block["cases"][3]["meta"]["frequency_text"]["axial"] == "15-20"
    assert len(block["cases"]) == 4
    excluded = block["excluded_observations"]
    assert [e["specimen_id"] for e in excluded] == ["M"]
    assert "frequenc" in excluded[0]["reason"]


def test_revised_missing_stress_flagged(tmp_path):
    empty = revised_row({"Specimen name": "E", "Amplitute AXIAL stress  (MPa)": None})
    block = _one(write_revised_workbook(tmp_path / "R.xlsx", _tension_rows() + [empty]))
    assert len(block["cases"]) == 3
    assert block["excluded_observations"][0]["reason"] == "missing stress component"


def test_legacy_null_stress_reads_zero(tmp_path):
    from workbook_factory import basquin_rows, write_workbook

    rows = basquin_rows()
    rows[0][5] = "Null"  # mean radial stress of a uniaxial row
    block = _one(write_workbook(tmp_path / "L.xlsx", rows=rows))
    assert block["source"]["format"] == "legacy"
    assert len(block["cases"]) == 15


def test_bending_and_amplitude_aliases():
    from multifat.imp.schema import DATASET_REQUIRED, find_col

    cols = [
        "Mean BENDING moment (Nmm)",
        "Amplitude BENDING stress (MPa)",
        "Mean BENDING stress (MPa)",
        "Amplitude AXIAL load (N)",
    ]
    assert find_col(cols, DATASET_REQUIRED["amp_ax"]) == "Amplitude BENDING stress (MPa)"
    assert find_col(cols, DATASET_REQUIRED["mean_ax"]) == "Mean BENDING stress (MPa)"
    with pytest.raises(KeyError):
        find_col(["Amplitude AXIAL load (N)"], DATASET_REQUIRED["amp_ax"])


def test_pressure_without_hoop_stress_flagged(tmp_path):
    row = revised_row({"Specimen name": "P", "Mean internal PRESSURE  (MPa)": 20})
    block = _one(write_revised_workbook(tmp_path / "Q.xlsx", _tension_rows() + [row]))
    assert len(block["cases"]) == 3
    assert block["excluded_observations"][0]["reason"] == (
        "internal pressure not converted to hoop stress"
    )


def test_identity_conflict_does_not_supply_fatigue_limits(tmp_path):
    rows = [
        ["r1", "A", 1992, "30NCD16", 195000, 0.3, 1160, 1020, 710, 450, 1e6, -1],
        ["r2", "B", 2024, "30NCD16", 195000, 0.3, 1160, 1020, 388, 450, 1e6, 0],
    ]
    path = write_revised_workbook(tmp_path / "2_CX18.xlsx", _tension_rows(), summary_rows=rows)
    block = _one(path)
    assert block["identity_issues"] and "30NCD16" in block["identity_issues"][0]
    assert "no compatible R = -1" in block["property_issues"]["Sigm1"]
    assert block["material"]["Tensile_fatigue_limit_MPa"] is None
    assert block["material"]["Torsional_fatigue_limit_MPa"] is None


def test_missing_nu_blocks_strain_models_only(tmp_path):
    from multifat.cli.pipeline import run_workbook
    from multifat.models.registry import stress_only_models

    stress_only = stress_only_models()
    assert {"CROSSLAND", "FIN", "MATAKE", "MCD"} <= stress_only and "SWT" not in stress_only
    rows = [["r", "A", 2020, "Synthetic steel", 206000, "Null", 900, 750, 420, 300, 2e6, -1]]
    path = write_revised_workbook(tmp_path / "N.xlsx", _tension_rows(), summary_rows=rows)
    block = _one(path)
    assert block["material"]["nu"] is None
    case = block["cases"][0]
    assert np.isfinite(case[1][1]["stress"]).all() and np.isnan(case[1][1]["strain"]).all()
    out = {
        o.model: o
        for o in run_workbook(
            path, models=["CROSSLAND", "SWT"], results_dir=tmp_path / "res", calibrate=False
        )
    }
    assert out["CROSSLAND"].status == "ok"
    assert out["SWT"].status == "skipped" and "needs E and nu" in out["SWT"].reason


@pytest.mark.skipif(not (ROOT / "Experimental_Data/6_42CrMo4_QT.xlsx").exists(), reason="data")
def test_all_active_workbooks_import():
    from multifat.cli.run import validate_workbook

    for path in sorted((ROOT / "Experimental_Data").glob("*.xlsx")):
        rep = validate_workbook(path)
        assert rep.material == material_id(path.stem)
        assert rep.ok or rep.errors, path.name


def test_mixed_ratio_summary_uses_reversed_source_for_models(tmp_path):
    from multifat.cli.pipeline import run_workbook

    rows = [
        ["r1", "A", 2018, "Synthetic steel", 206000, 0.3, 900, 750, 420, 300, 2e6, -1],
        ["r2", "B", 2024, "Synthetic steel", 206000, 0.3, 900, 750, 600, 300, 1e6, 0],
    ]
    path = write_revised_workbook(
        tmp_path / "Synthetic_steel.xlsx", _tension_rows(), summary_rows=rows
    )
    out = {
        o.model: o
        for o in run_workbook(
            path, models=["CROSSLAND", "SWT"], results_dir=tmp_path / "res", calibrate=False
        )
    }
    assert out["CROSSLAND"].status == "ok"
    assert out["SWT"].status == "ok"


@pytest.mark.skipif(not (ROOT / "Experimental_Data/14_42CrMo4_Reis.xlsx").exists(), reason="data")
def test_batch_workbook_counts():
    """Decisions d055, d056, d059: separate batches, AISI316L fixes, nu for four steels."""
    data = import_experimental_data_grouped_by_material(ROOT / "Experimental_Data")
    reis, berto, ss = data["42CrMo4_Reis"], data["Ti6Al4V_Berto"], data["AISI316L"]
    assert (len(reis["cases"]), len(reis["runouts"])) == (25, 4)
    assert reis["material"]["sigma_f"] == 1154 and reis["material"]["Tensile_strength_MPa"] == 1100
    assert (len(berto["cases"]), len(berto["runouts"])) == (48, 0)
    assert (len(ss["cases"]), len(ss["runouts"])) == (39, 13)
    jan90 = [c for c in list(ss["cases"].values()) + list(ss["runouts"].values())
             if c["meta"]["phase_tor_deg"] == 90]  # fmt: skip
    assert len(jan90) == 12
    assert (ss["material"]["E_MPa"], ss["material"]["nu"]) == (193000, 0.3)
    assert (
        ss["material"]["Tensile_fatigue_limit_MPa"],
        ss["material"]["Torsional_fatigue_limit_MPa"],
    ) == (241, 144)
    assert (
        data["Ti6Al4V"]["material"]["Tensile_fatigue_limit_MPa"],
        data["Ti6Al4V"]["material"]["Torsional_fatigue_limit_MPa"],
    ) == (500, 422.5)
    assert data["CX18"]["material"]["Tensile_fatigue_limit_MPa"] is None
    assert data["CX18"]["material"]["Torsional_fatigue_limit_MPa"] is None
    for m in ("CX18", "30NCD16", "25CrMo4", "982FA"):
        assert data[m]["material"]["nu"] == 0.3, m


def test_mixed_ratio_summary_reaches_every_entry_point(tmp_path):
    """Run and pipeline paths share the selected R = -1 value."""
    from multifat.cli.run import build_params, load_material

    rows = [
        ["r1", "A", 2018, "Synthetic steel", 206000, 0.3, 900, 750, 420, 300, 2e6, -1],
        ["r2", "B", 2024, "Synthetic steel", 206000, 0.3, 900, 750, 600, 300, 1e6, 0],
    ]
    path = write_revised_workbook(
        tmp_path / "Synthetic_steel.xlsx", _tension_rows(), summary_rows=rows
    )
    data, m = load_material(path)
    assert build_params(data[m]["material"])["Sigm1"] == 420
    assert "Sigm1" not in data[m]["property_issues"]
    assert data[m]["fatigue_limit_sources"]["Sigm1"][0]["excel_row"] == 2


def test_reversed_fatigue_limits_average_by_loading_mode(tmp_path):
    rows = [
        ["doi-a", "A", 2018, "Synthetic steel", 206000, 0.3, 900, 750, 420, 300, 2e6, -1],
        ["doi-b", "B", 2020, "Synthetic steel", 206000, 0.3, 900, 750, 440, 320, 2e6, -1],
        ["doi-c", "C", 2024, "Synthetic steel", 206000, 0.3, 900, 750, 600, 500, 1e6, 0],
    ]
    path = write_revised_workbook(
        tmp_path / "Synthetic_steel.xlsx", _tension_rows(), summary_rows=rows
    )
    block = _one(path)
    assert block["material"]["Tensile_fatigue_limit_MPa"] == 430
    assert block["material"]["Torsional_fatigue_limit_MPa"] == 310
    assert [source["reference"] for source in block["fatigue_limit_sources"]["Sigm1"]] == [
        "doi-a", "doi-b"
    ]
    assert [source["excel_row"] for source in block["fatigue_limit_sources"]["Taum1"]] == [
        2, 3
    ]


def test_no_reversed_torsional_source_leaves_property_unavailable(tmp_path):
    rows = [
        ["doi-a", "A", 2018, "Synthetic steel", 206000, 0.3, 900, 750, 420, None, 2e6, -1],
        ["doi-b", "B", 2024, "Synthetic steel", 206000, 0.3, 900, 750, None, 500, 1e6, 0],
    ]
    path = write_revised_workbook(
        tmp_path / "Synthetic_steel.xlsx", _tension_rows(), summary_rows=rows
    )
    block = _one(path)
    assert block["material"]["Tensile_fatigue_limit_MPa"] == 420
    assert block["material"]["Torsional_fatigue_limit_MPa"] is None
    assert "Taum1" in block["property_issues"]


def test_censored_source_is_not_an_exact_fatigue_limit(tmp_path):
    rows = [
        ["doi-a", "A", 2018, "Synthetic steel", 206000, 0.3, 900, 750, ">420", 300, 2e6, -1],
        ["doi-b", "B", 2020, "Synthetic steel", 206000, 0.3, 900, 750, 440, 320, 2e6, -1],
    ]
    path = write_revised_workbook(
        tmp_path / "Synthetic_steel.xlsx", _tension_rows(), summary_rows=rows
    )
    wb = load_workbook(path)
    wb["Summary"]["I1"] = "Tensile fatigue limit* (MPa)"
    wb.save(path)
    block = _one(path)
    assert block["material"]["Tensile_fatigue_limit_MPa"] == 440
    assert [source["excel_row"] for source in block["fatigue_limit_sources"]["Sigm1"]] == [3]


def test_average_cannot_replace_missing_source_limit(tmp_path):
    rows = [
        ["doi-a", "A", 2018, "Synthetic steel", 206000, 0.3, 900, 750, None, 300, 2e6, -1]
    ]
    path = write_revised_workbook(
        tmp_path / "Synthetic_steel.xlsx", _tension_rows(), summary_rows=rows
    )
    wb = load_workbook(path)
    wb["Summary"]["I3"] = 420
    wb.save(path)
    block = _one(path)
    assert block["material"]["Tensile_fatigue_limit_MPa"] is None
    assert "Sigm1" in block["property_issues"]


def test_blank_specimen_and_frequency_are_empty_text(tmp_path):
    """Minor: an empty cell is not the text 'nan'; no row number stands in for the specimen."""
    a = revised_row({"Specimen name": None, "Freq AXIAL load[Hz]": None})
    b = revised_row({"Specimen name": None, "Freq AXIAL load[Hz]": None})
    block = _one(write_revised_workbook(tmp_path / "B.xlsx", [a, b]))
    c0, c1 = block["cases"][0]["meta"], block["cases"][1]["meta"]
    assert c0["specimen_id"] == "" and c0["frequency_text"]["axial"] == ""
    from multifat.report.subsets import observation_index

    idx = observation_index(block)
    assert idx[("case", 0)]["Case_id"] == idx[("case", 1)]["Case_id"]  # same load, replicates


def test_excel_lock_files_ignored(tmp_path):
    """Minor: Excel's ~$ lock files are not workbooks."""
    from multifat.cli.run import resolve_workbook

    write_revised_workbook(tmp_path / "A.xlsx", _tension_rows())
    (tmp_path / "~$A.xlsx").write_bytes(b"lock")
    assert list(import_experimental_data_grouped_by_material(tmp_path)) == ["A"]
    assert resolve_workbook("A", tmp_path) == tmp_path / "A.xlsx"
